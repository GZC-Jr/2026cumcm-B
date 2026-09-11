"""
Convert Word documents to Markdown files.
"""
import sys
from pathlib import Path

try:
    import pypandoc
    USING_PYPANDOC = True
except ImportError:
    USING_PYPANDOC = False
    try:
        from docx import Document
        USING_DOCX = True
    except ImportError:
        USING_DOCX = False


def convert_with_pypandoc(docx_path: Path, md_path: Path):
    """Convert using pypandoc (requires pandoc installed)."""
    output = pypandoc.convert_file(str(docx_path), 'md', outputfile=str(md_path))
    return True


def convert_with_docx(docx_path: Path, md_path: Path):
    """Convert using python-docx (basic text extraction)."""
    doc = Document(docx_path)
    
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            # Basic heading detection
            if para.style.name.startswith('Heading'):
                level = para.style.name.replace('Heading', '').strip()
                if level.isdigit():
                    lines.append('#' * int(level) + ' ' + text)
                else:
                    lines.append('## ' + text)
            else:
                lines.append(text)
            lines.append('')
    
    # Handle tables
    for table in doc.tables:
        lines.append('')
        for i, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells]
            lines.append('| ' + ' | '.join(cells) + ' |')
            if i == 0:  # Add header separator
                lines.append('| ' + ' | '.join(['---'] * len(cells)) + ' |')
        lines.append('')
    
    md_path.write_text('\n'.join(lines), encoding='utf-8')
    return True


def convert_docx_to_md(docx_path: str) -> str:
    """
    Convert a .docx file to .md in the same directory.
    
    Args:
        docx_path: Path to the Word document
        
    Returns:
        Path to the created Markdown file
    """
    docx_path = Path(docx_path)
    
    if not docx_path.exists():
        raise FileNotFoundError(f"File not found: {docx_path}")
    
    if docx_path.suffix.lower() != '.docx':
        raise ValueError(f"Not a .docx file: {docx_path}")
    
    # Create .md path in same directory
    md_path = docx_path.with_suffix('.md')
    
    print(f"Converting: {docx_path.name}")
    print(f"Output: {md_path.name}")
    
    try:
        if USING_PYPANDOC:
            print("Using pypandoc converter...")
            convert_with_pypandoc(docx_path, md_path)
        elif USING_DOCX:
            print("Using python-docx converter (basic)...")
            convert_with_docx(docx_path, md_path)
        else:
            raise ImportError(
                "No conversion library available. Install either:\n"
                "  pip install pypandoc  (requires pandoc)\n"
                "  pip install python-docx  (basic conversion)"
            )
        
        print(f"[OK] Created: {md_path}")
        return str(md_path)
        
    except Exception as e:
        print(f"[ERROR] Error converting {docx_path.name}: {e}")
        raise


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Usage: python convert.py <docx_file1> [<docx_file2> ...]")
        sys.exit(1)
    
    for docx_file in sys.argv[1:]:
        try:
            convert_docx_to_md(docx_file)
        except Exception as e:
            print(f"Failed: {e}", file=sys.stderr)
            sys.exit(1)
    
    print(f"\nSuccessfully converted {len(sys.argv) - 1} file(s)")
